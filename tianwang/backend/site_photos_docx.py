#!/usr/bin/env python3
"""站点施工详情 Word 导出（按施工阶段分组，版式参照后台施工详情弹窗）"""
import io
import re
from datetime import datetime

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from daily_log import (
    _add_fullwidth_picture,
    _compact_paragraph,
    _printable_width,
    _set_cell,
    _set_run_font,
    _set_table_full_width,
)

STAGE_ORDER = ["勘察", "取电", "杆件", "设备", "光路"]
_INVALID_FNAME_CHARS = re.compile(r'[\\/:*?"<>|\r\n]+')


def _sanitize_filename(name: str, max_len: int = 48) -> str:
    text = _INVALID_FNAME_CHARS.sub("", (name or "").strip())
    text = re.sub(r"\s+", "", text)
    if not text:
        return "站点施工详情"
    return text[:max_len]


def build_site_photos_filename(site_name: str) -> str:
    return f"{_sanitize_filename(site_name)}-施工详情.docx"


def _add_heading(doc, text: str, size=14):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    _set_run_font(run, name="黑体", size=size, bold=True)


def _add_photo_caption(cell, content: str, person: str, upload_time: str):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _compact_paragraph(p, after_pt=0)
    cap = content or "施工照片"
    if upload_time:
        cap += f"\n{str(upload_time)[:16]}"
    if person:
        cap += f" · {person}"
    run = p.add_run(cap)
    _set_run_font(run, size=9)


def _add_stage_photos(doc, photos: list, img_width, max_bytes: int = 320_000):
    if not photos:
        p = doc.add_paragraph()
        run = p.add_run("暂无照片")
        _set_run_font(run, size=10.5)
        p.paragraph_format.space_after = Pt(8)
        return

    cols = 2
    half_w = int(img_width / cols - Cm(0.3))
    i = 0
    while i < len(photos):
        row_photos = photos[i : i + cols]
        tbl = doc.add_table(rows=2, cols=cols)
        tbl.style = "Table Grid"
        _set_table_full_width(tbl)
        for c in range(cols):
            img_cell = tbl.rows[0].cells[c]
            cap_cell = tbl.rows[1].cells[c]
            if c < len(row_photos):
                ph = row_photos[c]
                img_cell.text = ""
                p = img_cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                img_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                try:
                    if ph.get("photo_path"):
                        run = p.add_run()
                        from daily_log import _prepare_photo_stream

                        src = _prepare_photo_stream(ph["photo_path"], max_bytes=max_bytes)
                        run.add_picture(src, width=half_w)
                    else:
                        _set_cell(img_cell, "[照片缺失]", size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
                except Exception:
                    _set_cell(img_cell, "[照片加载失败]", size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
                _add_photo_caption(
                    cap_cell,
                    ph.get("content") or ph.get("build") or "",
                    ph.get("person") or "",
                    ph.get("uploadTime") or "",
                )
            else:
                img_cell.text = ""
                cap_cell.text = ""
        tbl.rows[0].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        doc.add_paragraph().paragraph_format.space_after = Pt(4)
        i += cols


def build_site_photos_docx(site: dict, photos: list, grouped: dict | None = None) -> bytes:
    """
    site: {name, area, status, createTime, photoCount, requiredCount, projectName}
    photos: [{content, build, stage, person, uploadTime, photo_path}, ...]
    grouped: {阶段名: [photos...]} 可选，不传则按 photo.stage 分组
    """
    site_name = (site.get("name") or "未命名站点").strip()
    doc = Document()
    doc.styles["Normal"].font.name = "宋体"
    doc.styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    img_width = _printable_width(section)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    title_run = title.add_run(f"{site_name} — 施工详情")
    _set_run_font(title_run, name="黑体", size=16, bold=True)

    info_tbl = doc.add_table(rows=2, cols=4)
    info_tbl.style = "Table Grid"
    _set_table_full_width(info_tbl)
    status = site.get("status") or "施工中"
    progress = f"{site.get('photoCount', 0)}/{site.get('requiredCount', 15)} 项"
    project = site.get("projectName") or site.get("projectId") or ""
    _set_cell(info_tbl.cell(0, 0), "地点", bold=True)
    _set_cell(info_tbl.cell(0, 1), site.get("area") or "-")
    _set_cell(info_tbl.cell(0, 2), "状态", bold=True)
    _set_cell(info_tbl.cell(0, 3), status)
    _set_cell(info_tbl.cell(1, 0), "项目", bold=True)
    _set_cell(info_tbl.cell(1, 1), project or "-")
    _set_cell(info_tbl.cell(1, 2), "施工进度", bold=True)
    _set_cell(info_tbl.cell(1, 3), progress)
    if site.get("createTime"):
        extra = doc.add_paragraph()
        extra_run = extra.add_run(f"创建时间：{str(site.get('createTime'))[:19]}")
        _set_run_font(extra_run, size=10)

    if grouped is None:
        grouped = {st: [] for st in STAGE_ORDER}
        grouped["其他"] = []
        for ph in photos:
            stage = ph.get("stage") or "其他"
            if stage not in grouped:
                grouped[stage] = []
            grouped[stage].append(ph)

    for stage in STAGE_ORDER + ["其他"]:
        stage_photos = grouped.get(stage) or []
        if stage == "其他" and not stage_photos:
            continue
        label = f"{stage}阶段" if stage != "其他" else "其他"
        _add_heading(doc, f"{label}（{len(stage_photos)}张）")
        _add_stage_photos(doc, stage_photos, img_width)

    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_run = footer.add_run(f"导出时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}")
    _set_run_font(footer_run, size=9)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
