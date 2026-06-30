#!/usr/bin/env python3
"""每日施工日志 Word 导出（版式参照项目参考模板）"""
import io
import json
import os
import re
from datetime import datetime

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from cos_storage import local_dir

LOG_META_FIELDS = ("engineeringName", "supervisorUnit", "ownerUnit", "contractorUnit")
_INVALID_FNAME_CHARS = re.compile(r'[\\/:*?"<>|\r\n]+')


def default_log_meta(project_key: str = "") -> dict:
    name = project_key if project_key and project_key != "__default__" else ""
    return {
        "engineeringName": name,
        "supervisorUnit": "",
        "ownerUnit": "中国联通南昌市分公司",
        "contractorUnit": "",
    }


def get_log_meta(conn, project_key: str = "") -> dict:
    key = project_key or "__default__"
    row = conn.execute("SELECT data_json FROM log_meta WHERE project_key=?", (key,)).fetchone()
    if not row:
        return default_log_meta(project_key)
    data = json.loads(row["data_json"])
    base = default_log_meta(project_key)
    base.update({k: data.get(k, base.get(k, "")) for k in LOG_META_FIELDS})
    return base


def save_log_meta(conn, project_key: str, data: dict) -> dict:
    key = project_key or "__default__"
    ts = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    item = default_log_meta(project_key)
    for k in LOG_META_FIELDS:
        if k in data:
            item[k] = (data.get(k) or "").strip()
    conn.execute(
        "INSERT INTO log_meta (project_key, data_json, updated_at) VALUES (?,?,?) "
        "ON CONFLICT(project_key) DO UPDATE SET data_json=excluded.data_json, updated_at=excluded.updated_at",
        (key, json.dumps(item, ensure_ascii=False), ts),
    )
    conn.commit()
    return item


def _set_run_font(run, name="宋体", size=10.5, bold=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def _compact_paragraph(p, *, after_pt=0):
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(after_pt)
    pf.line_spacing = 1.0


def _set_cell(cell, text, bold=False, size=10.5, align=WD_ALIGN_PARAGRAPH.LEFT, valign=WD_ALIGN_VERTICAL.CENTER):
    cell.text = ""
    cell.vertical_alignment = valign
    p = cell.paragraphs[0]
    p.alignment = align
    _compact_paragraph(p)
    run = p.add_run(str(text or ""))
    _set_run_font(run, size=size, bold=bold)


def _work_lines_font_size(lines) -> float:
    n = len(lines or [])
    if n > 10:
        return 9.5
    if n > 6:
        return 10.5
    if n > 4:
        return 11.0
    return 12.0


def _set_cell_numbered_lines(cell, lines, size=12):
    cell.text = ""
    cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    lines = [str(line).strip() for line in (lines or []) if str(line).strip()]
    if not lines:
        return
    for i, line in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        _compact_paragraph(p, after_pt=1 if i < len(lines) - 1 else 0)
        run = p.add_run(f"{i + 1}. {line}")
        _set_run_font(run, size=size)


def _add_page_break(doc):
    p = doc.add_paragraph()
    run = p.add_run()
    run.add_break(WD_BREAK.PAGE)


def _merge_row(table, row_idx, col_start, col_end):
    if col_start >= col_end:
        return
    start = table.cell(row_idx, col_start)
    for col in range(col_start + 1, col_end + 1):
        start.merge(table.cell(row_idx, col))


def _printable_width(section):
    return section.page_width - section.left_margin - section.right_margin


def _printable_height(section):
    return section.page_height - section.top_margin - section.bottom_margin


def _set_row_height(row, height, *, at_least=False):
    row.height = int(height)
    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST if at_least else WD_ROW_HEIGHT_RULE.EXACTLY


def _set_table_full_width(table):
    """表格横向铺满可打印区域。"""
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        tbl.insert(0, tbl_pr)
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "pct")
    tbl_w.set(qn("w:w"), "5000")
    if tbl_pr.find(qn("w:tblLayout")) is None:
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        tbl_pr.append(layout)
    table.autofit = False
    table.allow_autofit = False


def _set_table_rows_no_split(table):
    """禁止表格行跨页断开，避免第一页表被拆到第二页。"""
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:cantSplit")) is None:
            tr_pr.append(OxmlElement("w:cantSplit"))


def _layout_first_page_table(table, section):
    """第一页主表固定在一页内：标题下方至页底，行高精确分配。"""
    printable_h = _printable_height(section)
    title_reserve = Cm(1.25)
    available = printable_h - title_reserve

    header_h = Cm(0.64)
    body = available - header_h * 4
    row_heights = [
        header_h,
        header_h,
        header_h,
        header_h,
        body * 0.47,   # 工作情况
        body * 0.17,   # 存在问题
        body * 0.18,   # 解决结果
        body * 0.18,   # 备注
    ]
    for row, height in zip(table.rows, row_heights):
        _set_row_height(row, height, at_least=False)

    label_w = Cm(2.3)
    content_w = (_printable_width(section) - label_w) / 5
    for row in table.rows:
        row.cells[0].width = label_w
        for col in range(1, 6):
            row.cells[col].width = content_w

    _set_table_rows_no_split(table)


def _sanitize_filename_part(name: str, max_len: int = 36) -> str:
    text = _INVALID_FNAME_CHARS.sub("", (name or "").strip())
    text = re.sub(r"\s+", "", text)
    if not text:
        return "天网"
    if len(text) > max_len:
        return text[:max_len]
    return text


def build_export_filename(day: str, project: str = "", meta: dict | None = None) -> str:
    """文件名：项目名-施工日志-YYYYMMDD.docx"""
    meta = meta or {}
    dt = datetime.strptime(day, "%Y-%m-%d")
    date_tag = dt.strftime("%Y%m%d")
    eng = (meta.get("engineeringName") or "").strip()
    proj = (project or "").strip()
    if proj and (not eng or len(proj) <= len(eng)):
        base = proj
    elif eng:
        base = eng
    else:
        base = proj or "天网"
    return f"{_sanitize_filename_part(base)}-施工日志-{date_tag}.docx"


def _find_photo_path(meta: dict) -> str | None:
    fp = meta.get("filepath") or ""
    if fp and os.path.isfile(fp):
        return fp
    fname = meta.get("filename") or ""
    if fname:
        p = os.path.join(local_dir("uploads"), fname)
        if os.path.isfile(p):
            return p
    return None


def fetch_day_uploads(conn, day: str, project: str = "", person: str = ""):
    """day: YYYY-MM-DD"""
    rows = conn.execute(
        "SELECT id, filename, filepath, meta_json, created_at FROM uploads WHERE created_at LIKE ? ORDER BY created_at",
        (f"{day}%",),
    ).fetchall()
    items = []
    for r in rows:
        meta = json.loads(r["meta_json"] or "{}")
        if project and (meta.get("project") or "") != project:
            continue
        if person and (meta.get("person") or "") != person:
            continue
        photo_path = r["filepath"] if r["filepath"] and os.path.isfile(r["filepath"]) else None
        if not photo_path:
            photo_path = _find_photo_path({"filename": r["filename"], **meta})
        items.append(
            {
                "id": r["id"],
                "time": r["created_at"],
                "filename": r["filename"],
                "meta": meta,
                "photo_path": photo_path,
            }
        )
    return items


def _format_date_short(day: str) -> str:
    dt = datetime.strptime(day, "%Y-%m-%d")
    return f"{dt.year}-{dt.month}-{dt.day}"


def _build_work_lines(items: list) -> list[str]:
    lines = []
    seen = set()
    for it in items:
        m = it["meta"]
        point = (m.get("point") or "").strip()
        build = (m.get("build") or "").strip()
        key = f"{point}|{build}"
        if key in seen:
            continue
        seen.add(key)
        line = f"{point}{build}" if point and build else (point or build or "")
        extra = m.get("extraFields") or {}
        extra_vals = [str(v).strip() for v in extra.values() if v]
        if extra_vals:
            line += "".join(extra_vals)
        if line:
            lines.append(line)
    return lines


def _collect_site_persons(items: list, recorder: str = "") -> str:
    names = []
    for it in items:
        name = (it["meta"].get("person") or "").strip()
        if name and name not in names:
            names.append(name)
    if names:
        return "、".join(names)
    return recorder or ""


def _photo_caption(meta: dict) -> str:
    point = (meta.get("point") or "").strip()
    build = (meta.get("build") or "").strip()
    if point:
        return point
    return build or "施工照片"


def _prepare_photo_stream(photo_path: str, max_bytes: int = 280_000):
    """压缩/缩小照片后再嵌入 Word，避免大图导致生成过慢或手机端下载超时。"""
    try:
        from image_compress import compress_image_bytes

        with open(photo_path, "rb") as f:
            data = f.read()
        out, _ = compress_image_bytes(data, max_size=max_bytes)
        return io.BytesIO(out)
    except Exception:
        return photo_path


def _add_fullwidth_picture(cell, photo_path: str, width, max_bytes: int = 280_000):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    src = _prepare_photo_stream(photo_path, max_bytes=max_bytes)
    run.add_picture(src, width=width)


PHOTO_EXPORT_PRESETS = {
    "none": {"max_photos": 0, "max_bytes": 0},
    "share": {"max_photos": 12, "max_bytes": 280_000},
    "full": {"max_photos": 40, "max_bytes": 500_000},
}


def build_daily_log_docx(
    day: str,
    project: str,
    items: list,
    recorder: str = "",
    meta: dict | None = None,
    photo_mode: str = "share",
    photo_ids: set | None = None,
) -> bytes:
    meta = meta or {}
    eng_name = meta.get("engineeringName") or project or "天网项目"
    supervisor = meta.get("supervisorUnit") or ""
    owner = meta.get("ownerUnit") or ""
    contractor = meta.get("contractorUnit") or ""
    date_short = _format_date_short(day)
    work_lines = _build_work_lines(items)
    site_persons = _collect_site_persons(items, recorder)

    doc = Document()
    doc.styles["Normal"].font.name = "宋体"
    doc.styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)
    section.top_margin = Cm(1.27)
    section.bottom_margin = Cm(1.27)
    img_width = _printable_width(section)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(4)
    title.paragraph_format.line_spacing = 1.0
    title.paragraph_format.keep_with_next = True
    title_run = title.add_run("施工日志")
    _set_run_font(title_run, name="黑体", size=18, bold=True)

    tbl = doc.add_table(rows=8, cols=6)
    tbl.style = "Table Grid"
    _set_table_full_width(tbl)

    _set_cell(tbl.cell(0, 0), "工程名称", bold=True)
    _set_cell(tbl.cell(0, 1), eng_name)
    _merge_row(tbl, 0, 1, 5)

    _set_cell(tbl.cell(1, 0), "建设单位", bold=True)
    _set_cell(tbl.cell(1, 1), owner)
    _merge_row(tbl, 1, 1, 5)

    _set_cell(tbl.cell(2, 0), "监理单位", bold=True)
    _set_cell(tbl.cell(2, 1), supervisor)
    _merge_row(tbl, 2, 1, 2)
    _set_cell(tbl.cell(2, 3), "承建单位", bold=True)
    _set_cell(tbl.cell(2, 4), contractor)
    _merge_row(tbl, 2, 4, 5)

    _set_cell(tbl.cell(3, 0), "记录日期", bold=True)
    _set_cell(tbl.cell(3, 1), date_short)
    _merge_row(tbl, 3, 1, 2)
    _set_cell(tbl.cell(3, 3), "现场人员：", bold=True)
    _merge_row(tbl, 3, 3, 4)
    _set_cell(tbl.cell(3, 5), site_persons)

    _set_cell(tbl.cell(4, 0), "工 作\n情 况", bold=True)
    _merge_row(tbl, 4, 0, 1)
    work_font = _work_lines_font_size(work_lines)
    _set_cell_numbered_lines(tbl.cell(4, 2), work_lines, size=work_font)
    _merge_row(tbl, 4, 2, 5)

    _set_cell(tbl.cell(5, 0), "存 在\n问 题", bold=True)
    _merge_row(tbl, 5, 0, 1)
    _set_cell(tbl.cell(5, 2), "", valign=WD_ALIGN_VERTICAL.TOP)
    _merge_row(tbl, 5, 2, 5)

    _set_cell(tbl.cell(6, 0), "解 决\n结 果", bold=True)
    _merge_row(tbl, 6, 0, 1)
    _set_cell(tbl.cell(6, 2), "", valign=WD_ALIGN_VERTICAL.TOP)
    _merge_row(tbl, 6, 2, 5)

    _set_cell(
        tbl.cell(7, 0),
        "备注：\n\n注：此表一式三份，建设单位、监理单位、承建单位各执一份。",
        size=10,
        valign=WD_ALIGN_VERTICAL.TOP,
    )
    _merge_row(tbl, 7, 0, 5)

    _layout_first_page_table(tbl, section)

    photo_items = [it for it in items if it.get("photo_path")]
    if photo_ids is not None:
        photo_items = [it for it in photo_items if str(it["id"]) in photo_ids]
    preset = PHOTO_EXPORT_PRESETS.get(photo_mode) or PHOTO_EXPORT_PRESETS["share"]
    max_photos = preset["max_photos"]
    max_bytes = preset["max_bytes"]
    skipped = 0
    manual_pick = photo_ids is not None
    if max_photos <= 0:
        photo_items = []
    elif not manual_pick and len(photo_items) > max_photos:
        skipped = len(photo_items) - max_photos
        photo_items = photo_items[:max_photos]
    elif manual_pick and len(photo_items) > 60:
        skipped = len(photo_items) - 60
        photo_items = photo_items[:60]
    if photo_items:
        _add_page_break(doc)
        attach = doc.add_paragraph()
        attach_run = attach.add_run("附件：施工照片")
        _set_run_font(attach_run, size=12, bold=True)
        if skipped:
            note_p = doc.add_paragraph()
            note_run = note_p.add_run(f"（另有 {skipped} 张照片未嵌入，请至平台照片管理查看）")
            _set_run_font(note_run, size=10.5)

        for idx, it in enumerate(photo_items, 1):
            if idx > 1:
                _add_page_break(doc)
            caption = _photo_caption(it["meta"])
            cap = doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap_run = cap.add_run(f"{idx}. {caption}")
            _set_run_font(cap_run, name="黑体", size=16)

            img_tbl = doc.add_table(rows=1, cols=1)
            img_tbl.style = "Table Grid"
            cell = img_tbl.rows[0].cells[0]
            try:
                _add_fullwidth_picture(cell, it["photo_path"], img_width, max_bytes=max_bytes)
            except Exception:
                _set_cell(cell, f"[照片：{it['filename']}]", valign=WD_ALIGN_VERTICAL.CENTER)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
